"""PPTX generator — KPMG-style deck template. Sections already contain resolved
data (tables / chart PNG bytes or chart specs / insight text) — the generator only
lays them out; it never receives a manifest key or invents a number.

Look (unofficial KPMG-style homage): white canvas, left-aligned, kicker -> headline ->
hairline header, hairline-only tables, blue-ladder charts, solid-navy dividers,
footer with classification + page number. Drop an official master at
templates/kpmg_master.pptx and its slide master/background is used instead of the
python-pptx default.

Section kinds (spec["sections"]):
  table | chart | narrative                       — generic reports (create_report etc.)
  agenda | divider | kpis | key_values | timeline
  chart_insights | two_charts | cards             — deep-dive decks
"""
import io
import re
from datetime import date, datetime
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import charts, theme

TEMPLATE = Path(__file__).parent / "templates" / "kpmg_master.pptx"

W_IN, H_IN = 13.333, 7.5
MX = 0.6                                  # side margin
CX, CY, CW, CH = MX, 1.55, W_IN - 2 * MX, 5.3   # content box (below header, above footer)
FONT = theme.FONT_FAMILY
ROWS_PER_SLIDE, MAX_COLS = 13, 8

_ACRONYMS = {"p2p", "po", "pr", "grn", "gst", "gstin", "msme", "pan", "id", "ytd", "otd", "fi", "mm", "capex", "opex", "uom"}
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")
LADDER = [theme.NAVY, theme.MED, theme.LIGHT, theme.TEAL_POS]


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


NAVY, LIGHT, WHITE = _rgb(theme.NAVY), _rgb(theme.LIGHT), _rgb("#FFFFFF")
BODY, CAPTION = _rgb(theme.BODY), _rgb(theme.CAPTION)
HAIR, PANEL, CHIP = _rgb(theme.HAIRLINE), _rgb(theme.PANEL), _rgb(theme.CHIP)
RED, SOFT_BLUE = _rgb(theme.RED), _rgb("#B0C4DE")


# ── value formatting ─────────────────────────────────────────────────────────────

def pretty_header(col: str) -> str:
    return " ".join(w.upper() if w.lower() in _ACRONYMS else w.capitalize() for w in str(col).split("_"))


def fmt_cell(col: str, v, money: bool = False, plain: bool = False) -> tuple[str, bool]:
    """Display string for a table cell + whether it is numeric (right-aligned)."""
    if v is None or v == "":
        return "—", False
    if isinstance(v, bool):
        return ("Yes" if v else "No"), False
    if isinstance(v, (date, datetime)):
        return v.strftime("%d-%b-%Y"), False
    if isinstance(v, (dict, list, tuple)):
        return f"{len(v)} item{'s' if len(v) != 1 else ''}", False
    unit = "inr" if money else "number" if plain else charts._unit_for(col)   # plain: no money/% guessing from the name
    if isinstance(v, str):
        if _DATE_RE.match(v):
            try:
                return datetime.strptime(v[:10], "%Y-%m-%d").strftime("%d-%b-%Y"), False
            except ValueError:
                pass
        if unit == "inr" and _NUM_RE.match(v) and "." in v:   # TEXT money columns, e.g. '111602923.21'
            return charts._fmt(float(v), "inr"), True
        return (v if len(v) <= 70 else v[:69] + "…"), False
    if isinstance(v, (int, float)):
        if unit in ("inr", "pct"):
            return charts._fmt(float(v), unit), True
        return (f"{v:,}" if isinstance(v, int) else f"{v:,.1f}".rstrip("0").rstrip(".")), True
    return str(v)[:40], False


# ── low-level drawing helpers ────────────────────────────────────────────────────

def _rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.line.fill.background()
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    style = sh._element.find(qn("p:style"))   # theme effectRef = the drop shadow LibreOffice/PowerPoint draw
    if style is not None:
        sh._element.remove(style)
    return sh


def _runs(p, text, size, color, bold=False):
    """Add runs to paragraph p, honouring **bold** markers."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size, r.font.name = Pt(size), FONT
        r.font.bold = bold or (i % 2 == 1)
        r.font.color.rgb = color


def _text(slide, text, x, y, w, h, size=12, bold=False, color=BODY, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _runs(p, line, size, color, bold)
    return box


def _fit(items: list[str], w: float, h: float, base: int, min_size: int = 9) -> tuple[list[str], int]:
    """Largest font size (<= base) at which the bullets fit the box; drops trailing bullets if even min_size overflows."""
    items = list(items)
    while True:
        for size in range(base, min_size - 1, -1):
            cpl = max(int(w / (size * 0.0072)), 10)          # ~chars per line for Arial
            need = sum(-(-len(i) // cpl) for i in items) * size * 1.22 / 72 + len(items) * 6 / 72
            if need <= h:
                return items, size
        if len(items) <= 1:
            return items, min_size
        items = items[:-1]


def _bullets(slide, items, x, y, w, h, size=13, color=BODY, bullet=True):
    items, size = _fit(items, w - (0.25 if bullet else 0), h, size)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "228600")
            pPr.set("indent", "-228600")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")   # must follow spcAft
        _runs(p, item, size, color)
    return box


def parse_bullets(text: str) -> list[str]:
    """Narrative text -> list of lines; '- '/'* '/'• ' markers stripped, blank lines dropped."""
    out = []
    for line in str(text).splitlines():
        line = re.sub(r"^\s*([-*•]|\d+[.)])\s+", "", line).strip()
        if line:
            out.append(line)
    return out


# ── cell borders (python-pptx has no API): bottom hairline only ──────────────────

def _cell_style(cell, fill, bottom=None):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):   # inserted at 0 in reverse -> L,R,T,B order
        ln = etree.Element(qn(tag))
        if tag == "a:lnB" and bottom is not None:
            ln.set("w", "6350")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", bottom)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(0, ln)
    cell.margin_left = cell.margin_right = Inches(0.07)
    cell.margin_top = cell.margin_bottom = Inches(0.025)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


# ── the deck ─────────────────────────────────────────────────────────────────────

class Deck:
    def __init__(self):
        if TEMPLATE.exists():
            self.prs = Presentation(str(TEMPLATE))
            ids = self.prs.slides._sldIdLst
            for sld in list(ids):          # drop the template's sample slides
                self.prs.part.drop_rel(sld.rId)
                ids.remove(sld)
            self.layout = next((l for l in self.prs.slide_layouts if l.name.lower() == "blank"),
                               self.prs.slide_layouts[len(self.prs.slide_layouts) - 1])
        else:
            self.prs = Presentation()
            self.layout = self.prs.slide_layouts[6]
        self.prs.slide_width, self.prs.slide_height = Inches(W_IN), Inches(H_IN)
        self.today = datetime.now().strftime("%d %b %Y")

    def _slide(self, bg=None):
        s = self.prs.slides.add_slide(self.layout)
        if bg is not None:
            _rect(s, 0, 0, W_IN, H_IN, bg)
        return s

    # ---- structural slides
    def cover(self, title, subtitle="", kicker="KPMG IntelliSource · Procurement Intelligence"):
        s = self._slide()
        _rect(s, 0, 0, W_IN, 0.1, NAVY)
        _text(s, kicker.upper(), MX, 1.7, 11, 0.3, size=11, bold=True, color=LIGHT)
        _text(s, title, MX, 2.1, 11.5, 1.9, size=40 if len(title) <= 55 else 30, bold=True, color=NAVY)
        if subtitle:
            _text(s, subtitle, MX, 4.2, 11, 0.9, size=16, color=CAPTION)
        _rect(s, MX, 6.5, W_IN - 2 * MX, 0.012, HAIR)
        _text(s, self.today, MX, 6.65, 3, 0.3, size=10, color=CAPTION)
        _text(s, "Document classification: Confidential", 4.0, 6.65, 5, 0.3, size=10, color=CAPTION)
        blk = _rect(s, W_IN - MX - 1.1, 6.62, 1.1, 0.42, NAVY)
        _text(s, "KPMG", W_IN - MX - 1.1, 6.62, 1.1, 0.42, size=14, bold=True, color=WHITE,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def divider(self, number, title, subtitle=""):
        s = self._slide(NAVY)
        _text(s, f"{number:02d}" if isinstance(number, int) else str(number), MX, 2.55, 3, 0.5,
              size=20, bold=True, color=LIGHT)
        _rect(s, MX, 3.1, 1.2, 0.04, LIGHT)
        _text(s, title, MX, 3.3, 11.5, 1.3, size=44, bold=True, color=WHITE)
        if subtitle:
            _text(s, subtitle, MX, 4.8, 11, 0.8, size=16, color=SOFT_BLUE)

    def closing(self, note=""):
        s = self._slide(NAVY)
        _rect(s, MX, 3.0, 1.2, 0.04, LIGHT)
        _text(s, "Thank you", MX, 3.2, 11, 1.0, size=44, bold=True, color=WHITE)
        _text(s, f"Generated by KPMG IntelliSource — Ask · {self.today}", MX, 4.5, 11, 0.4, size=14, color=SOFT_BLUE)
        _text(s, note or "All figures are drawn directly from the procurement data warehouse (read-only analysis).",
              MX, 5.0, 11, 0.6, size=11, color=SOFT_BLUE)

    def content(self, kicker, headline):
        """Content slide with kicker / headline / hairline header and footer."""
        s = self._slide()
        _text(s, (kicker or "").upper(), MX, 0.35, 11, 0.25, size=10, bold=True, color=LIGHT)
        _text(s, headline, MX, 0.62, W_IN - 2 * MX, 0.6, size=24 if len(headline) <= 60 else 18,
              bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)
        _rect(s, MX, 1.28, W_IN - 2 * MX, 0.012, HAIR)
        _rect(s, MX, 7.0, W_IN - 2 * MX, 0.012, HAIR)
        _text(s, "KPMG IntelliSource · Document classification: Confidential", MX, 7.08, 8, 0.25, size=9, color=CAPTION)
        _text(s, str(len(self.prs.slides)), W_IN - MX - 1, 7.08, 1, 0.25, size=9, color=CAPTION, align=PP_ALIGN.RIGHT)
        return s

    # ---- content blocks
    def kpi_cards(self, s, items, x, y, w, h):
        per_row = len(items) if len(items) <= 4 else 3
        rows = -(-len(items) // per_row)
        gap = 0.2
        cw, ch = (w - gap * (per_row - 1)) / per_row, (h - gap * (rows - 1)) / rows
        ch = min(ch, 1.3)
        for i, item in enumerate(items):
            label, value = item[0], str(item[1])
            cx, cy = x + (i % per_row) * (cw + gap), y + (i // per_row) * (ch + gap)
            _rect(s, cx, cy, cw, ch, CHIP)
            _rect(s, cx, cy, 0.05, ch, NAVY)
            _text(s, value, cx + 0.2, cy + 0.14, cw - 0.3, ch * 0.5, size=26 if len(value) <= 9 else 18 if len(value) <= 16 else 14,
                  bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
            _text(s, label, cx + 0.2, cy + ch * 0.62, cw - 0.3, ch * 0.34, size=10, color=CAPTION)
        return rows * (ch + gap) - gap

    def table(self, s, rows, x, y, w, h, cols=None, money_cols=(), plain=False):
        cols = cols or list(rows[0].keys())[:MAX_COLS]
        cells = [[fmt_cell(c, r.get(c), c in money_cols, plain) for c in cols] for r in rows]
        # column widths: proportional to the longest display string (header or cell), clamped
        lens = [max([len(pretty_header(c))] + [len(cells[i][j][0]) for i in range(len(rows))]) for j, c in enumerate(cols)]
        lens = [min(max(n, 6), 45) for n in lens]
        widths = [w * n / sum(lens) for n in lens]
        row_h = min(0.36, h / (len(rows) + 1))
        shape = s.shapes.add_table(len(rows) + 1, len(cols), Inches(x), Inches(y), Inches(w), Inches(row_h * (len(rows) + 1)))
        tbl = shape.table
        tbl.horz_banding = False
        tbl.first_row = True
        for j, cw in enumerate(widths):
            tbl.columns[j].width = Inches(cw)
        for i in range(len(rows) + 1):
            tbl.rows[i].height = Inches(row_h)
        size = 10 if len(cols) <= 6 else 9
        for j, c in enumerate(cols):
            numeric = any(cells[i][j][1] for i in range(len(rows)))
            cell = tbl.cell(0, j)
            _cell_style(cell, NAVY)
            cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
            _runs(cell.text_frame.paragraphs[0], pretty_header(c), size, WHITE, bold=True)
            for i in range(len(rows)):
                cell = tbl.cell(i + 1, j)
                _cell_style(cell, WHITE, bottom=theme.HAIRLINE.lstrip("#"))
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
                _runs(p, cells[i][j][0], size, BODY)
        return shape

    def chart(self, s, png, x, y, w, h):
        """Place PNG bytes inside the box keeping aspect ratio (never overflows)."""
        from PIL import Image
        iw, ih = Image.open(io.BytesIO(png)).size
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        s.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(pw), Inches(ph))

    def chart_spec(self, s, spec, x, y, w, h):
        self.chart(s, charts.render_png(spec, size=(w, h), title=False), x, y, w, h)

    def insight_panel(self, s, bullets, x, y, w, h, title="Key takeaways"):
        _rect(s, x, y, w, h, PANEL)
        _rect(s, x, y, w, 0.05, NAVY)
        _text(s, title.upper(), x + 0.2, y + 0.2, w - 0.4, 0.3, size=10, bold=True, color=NAVY)
        bullets = bullets[:6]
        _bullets(s, bullets, x + 0.2, y + 0.6, w - 0.4, h - 0.8, size=12)

    def key_values(self, s, pairs, x, y, w, h):
        _rect(s, x, y, w, h, PANEL)
        cols = 2 if len(pairs) > 13 else 1
        per = -(-len(pairs) // cols)
        colw = (w - 0.4) / cols
        row_h = min(0.42, (h - 0.3) / per)
        for i, (k, v) in enumerate(pairs):
            cx, cy = x + 0.2 + (i // per) * colw, y + 0.15 + (i % per) * row_h
            _text(s, k, cx, cy, colw * 0.4, row_h, size=10, color=CAPTION, anchor=MSO_ANCHOR.MIDDLE)
            _text(s, str(v), cx + colw * 0.42, cy, colw * 0.55, row_h, size=11, bold=True, color=BODY, anchor=MSO_ANCHOR.MIDDLE)

    def timeline(self, s, milestones, x, y, w, h):
        """Horizontal milestone line. milestone = {label, date, gap_days?, late?}."""
        n = len(milestones)
        step = w / n
        line_y = y + h * 0.5
        _rect(s, x + step / 2, line_y - 0.015, step * (n - 1), 0.03, HAIR)
        for i, m in enumerate(milestones):
            cx = x + step * i + step / 2
            _rect(s, cx - 0.11, line_y - 0.11, 0.22, 0.22, NAVY)
            _text(s, m["label"], cx - step / 2 + 0.05, line_y - 1.15, step - 0.1, 0.7, size=10, bold=True,
                  color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
            _text(s, m.get("date") or "—", cx - step / 2 + 0.05, line_y + 0.25, step - 0.1, 0.3, size=9,
                  color=CAPTION, align=PP_ALIGN.CENTER)
            if i and m.get("gap_days") is not None:
                late = m.get("late")
                _text(s, f"+{m['gap_days']} d", cx - step, line_y - 0.32, step, 0.2, size=9, bold=True,
                      color=RED if late else LIGHT, align=PP_ALIGN.CENTER)

    def cards(self, s, items, x, y, w, h):
        """Equal flat rectangles with 4 px top edges stepping down the blue ladder."""
        n = len(items)
        per_row = n if n <= 3 else 2
        rows = -(-n // per_row)
        gap = 0.25
        cw, ch = (w - gap * (per_row - 1)) / per_row, min((h - gap * (rows - 1)) / rows, 2.6)
        for i, (title, body) in enumerate(items):
            cx, cy = x + (i % per_row) * (cw + gap), y + (i // per_row) * (ch + gap)
            _rect(s, cx, cy, cw, ch, WHITE).line.color.rgb = HAIR
            _rect(s, cx, cy, cw, 0.055, _rgb(LADDER[i % len(LADDER)]))
            _text(s, title, cx + 0.2, cy + 0.2, cw - 0.4, 0.5, size=13, bold=True, color=NAVY)
            _text(s, body, cx + 0.2, cy + 0.75, cw - 0.4, ch - 0.9, size=11, color=BODY)

    # ---- section dispatch
    def section(self, sec):
        kind = sec.get("kind", "table")
        head, kick = sec.get("heading", ""), sec.get("kicker", "")
        if kind == "divider":
            return self.divider(sec.get("number", ""), head, sec.get("subtitle", ""))
        if kind == "agenda":
            s = self.content("Agenda", head or "What this deck covers")
            for i, (title, desc) in enumerate(sec["items"][:6]):
                yy = CY + 0.1 + i * 0.85
                _text(s, f"{i + 1:02d}", CX, yy, 1, 0.5, size=22, bold=True, color=LIGHT)
                _text(s, title, CX + 1.0, yy, 10, 0.4, size=18, bold=True, color=NAVY)
                _text(s, desc, CX + 1.0, yy + 0.4, 10.5, 0.35, size=12, color=CAPTION)
        elif kind == "kpis":
            s = self.content(kick or "Summary", head)
            used = self.kpi_cards(s, sec["kpis"], CX, CY, CW, 2.9 if len(sec["kpis"]) > 4 else 1.4)
            bl = sec.get("bullets") or []
            if bl:
                _text(s, sec.get("bullets_title", "Key takeaways").upper(), CX, CY + used + 0.3, CW, 0.3, size=10, bold=True, color=NAVY)
                _bullets(s, bl[:6], CX, CY + used + 0.7, CW, CH - used - 0.8, size=13)
        elif kind == "key_values":
            s = self.content(kick or "Overview", head)
            rows = sec.get("rows") or []
            kw = 4.9 if rows else CW
            self.key_values(s, sec["pairs"], CX, CY, kw, CH)
            if rows:
                self.table(s, rows[:ROWS_PER_SLIDE], CX + kw + 0.3, CY, CW - kw - 0.3, CH, cols=sec.get("cols"), money_cols=sec.get("money_cols", ()))
        elif kind == "timeline":
            s = self.content(kick or "Lifecycle", head)
            self.timeline(s, sec["milestones"], CX, CY, CW, 2.5)
            below_y = CY + 2.7
            if sec.get("chart"):
                self.chart_spec(s, sec["chart"], CX, below_y, 7.4, CH - 2.7)
                self.insight_panel(s, sec.get("insights", []), CX + 7.7, below_y, CW - 7.7, CH - 2.7)
            elif sec.get("insights"):
                self.insight_panel(s, sec["insights"], CX, below_y, CW, CH - 2.7)
        elif kind == "chart_insights":
            s = self.content(kick or "Analysis", head)
            if sec.get("insights"):
                self.chart_spec(s, sec["chart"], CX, CY, 7.7, CH)
                self.insight_panel(s, sec["insights"], CX + 8.0, CY, CW - 8.0, CH)
            else:
                self.chart_spec(s, sec["chart"], CX, CY, CW, CH)
        elif kind == "two_charts":
            s = self.content(kick or "Analysis", head)
            bl = sec.get("insights") or []
            ch = CH - (1.4 if bl else 0)
            cw = (CW - 0.4) / 2
            for i, spec in enumerate(sec["charts"][:2]):
                xx = CX + i * (cw + 0.4)
                _text(s, spec.get("title", ""), xx, CY, cw, 0.3, size=11, bold=True, color=NAVY)
                self.chart_spec(s, spec, xx, CY + 0.35, cw, ch - 0.35)
            if bl:
                _bullets(s, bl[:3], CX, CY + ch + 0.15, CW, 1.2, size=12)
        elif kind == "cards":
            s = self.content(kick or "Recommendations", head)
            self.cards(s, sec["items"][:6], CX, CY, CW, CH)
        elif kind == "narrative":
            s = self.content(kick or "Summary", head)
            items = parse_bullets(sec.get("narrative", ""))
            _bullets(s, items, CX, CY, CW, CH, size=14, bullet=len(items) > 1)
        elif kind == "chart" and sec.get("chart_png"):
            s = self.content(kick or "Chart", head)
            ins = sec.get("insights")
            self.chart(s, sec["chart_png"], CX, CY, 8.0 if ins else CW, CH)
            if ins:
                self.insight_panel(s, ins, CX + 8.3, CY, CW - 8.3, CH)
        else:  # table, paginated
            rows = sec.get("data") or []
            if not rows:
                s = self.content(kick or "Data", head)
                return _text(s, "No data.", CX, CY, 8, 0.4, size=12, color=CAPTION)
            cols = sec.get("cols") or list(rows[0].keys())[:MAX_COLS]
            pages = [rows[i:i + ROWS_PER_SLIDE] for i in range(0, len(rows), ROWS_PER_SLIDE)]
            for n, chunk in enumerate(pages):
                s = self.content(kick or "Data", head if n == 0 else f"{head} (cont.)")
                self.table(s, chunk, CX, CY, CW, CH - 0.3, cols=cols, money_cols=sec.get("money_cols", ()), plain=sec.get("plain", False))
                if len(pages) > 1:
                    _text(s, f"Rows {n * ROWS_PER_SLIDE + 1}–{n * ROWS_PER_SLIDE + len(chunk)} of {len(rows)}",
                          CX, CY + CH - 0.2, 6, 0.25, size=9, color=CAPTION)

    def save(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def build_pptx_report(spec: dict) -> bytes:
    """spec = {title, subtitle?, kicker?, closing_note?, sections: [{heading, kind, ...}]}"""
    deck = Deck()
    deck.cover(spec["title"], spec.get("subtitle", ""), spec.get("kicker", "KPMG IntelliSource · Ask"))
    for section in spec["sections"]:
        deck.section(section)
    deck.closing(spec.get("closing_note", ""))
    return deck.save()


if __name__ == "__main__":
    bar = {"chart_type": "bar", "title": "Spend", "x": "v", "y": "spend", "unit": "inr",
           "points": [{"x": f"Vendor {i}", "y": 1e7 * (i + 1)} for i in range(6)], "total_points": 6}
    grouped = {"chart_type": "grouped", "title": "Vendor vs portfolio", "unit": "number", "labels": ["PO→GRN", "GRN→Inv"],
               "series": {"Vendor": [44.4, 9.4], "Portfolio": [39.9, 8.0]}}
    big = [{"po": f"20000{i:05d}", "vendor": "L&T Construction Ltd", "net_order_value": f"{110000000 + i}.21",
            "document_date": "2023-01-18", "nested": [1, 2]} for i in range(30)]
    spec = {"title": "Self-check deck: a deliberately long title to exercise the shrink path", "subtitle": "sub", "sections": [
        {"kind": "agenda", "heading": "Agenda", "items": [("One", "first"), ("Two", "second")]},
        {"kind": "divider", "number": 1, "heading": "Section", "subtitle": "sub"},
        {"kind": "kpis", "heading": "KPIs", "kpis": [("PO value", "₹11.16 Cr"), ("Lines", "1"), ("Delay", "3 d"), ("Risk", "Medium"), ("x", "1"), ("y", "2")],
         "bullets": ["**Bold** point one", "Point two"]},
        {"kind": "key_values", "heading": "KV", "pairs": [("Vendor", "L&T"), ("City", "Mumbai")] * 5, "rows": big[:3]},
        {"kind": "timeline", "heading": "TL", "milestones": [
            {"label": "PR created", "date": "15-Jan-2023"}, {"label": "PO", "date": "18-Jan-2023", "gap_days": 3},
            {"label": "GRN", "date": "24-Feb-2023", "gap_days": 37, "late": True}, {"label": "Paid", "date": "07-May-2023", "gap_days": 72}],
         "chart": bar, "insights": ["Delivered 3 days late"]},
        {"kind": "chart_insights", "heading": "CI", "chart": bar, "insights": ["a", "b"]},
        {"kind": "two_charts", "heading": "TC", "charts": [bar, grouped], "insights": ["one"]},
        {"kind": "cards", "heading": "Cards", "items": [("A", "alpha body"), ("B", "beta body"), ("C", "gamma body")]},
        {"kind": "table", "heading": "Big table", "data": big},
        {"kind": "narrative", "heading": "Narr", "narrative": "- one\n- two **bold**"},
        {"kind": "chart", "heading": "PNG", "chart_png": charts.render_png(bar, size=(7.5, 4.5))},
    ]}
    data = build_pptx_report(spec)
    import os
    if os.environ.get("PPTX_OUT"):   # PPTX_OUT=/tmp/x.pptx python -m services.reports.pptx  -> inspect visually
        Path(os.environ["PPTX_OUT"]).write_bytes(data)
    prs = Presentation(io.BytesIO(data))
    W, H = prs.slide_width, prs.slide_height
    for n, sl in enumerate(prs.slides, 1):
        for sh in sl.shapes:
            assert sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= W + 10 and sh.top + sh.height <= H + 10, \
                f"slide {n}: {sh.name} out of bounds ({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f},{Emu(sh.width).inches:.2f},{Emu(sh.height).inches:.2f})"
    titles = [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]
    assert any("(cont.)" in t for t in titles), "long table was not paginated"
    assert not any("[{" in t or "nested" == t for t in titles), "raw repr leaked into a slide"
    print("pptx OK —", len(prs.slides), "slides,", len(data), "bytes")
