"""Excel report generator — openpyxl styled workbook. Takes a ReportSpec whose
sections already contain resolved row data (see report_tools.py) — never a
manifest key, never a model-authored number.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import theme

HEADER_FILL = PatternFill(start_color="00338D", end_color="00338D", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="00338D", bold=True, size=16)


def build_excel_report(spec: dict) -> bytes:
    """spec = {title, subtitle?, sections: [{heading, data: [dict,...]}, ...]}"""
    wb = Workbook()
    wb.remove(wb.active)

    for section in spec["sections"]:
        rows = section.get("data") or []
        heading = section["heading"][:31]  # Excel sheet name limit
        ws = wb.create_sheet(title=heading or "Sheet")

        ws["A1"] = section["heading"]
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:D1")

        if not rows:
            ws["A3"] = "No data."
            continue

        columns = list(rows[0].keys())
        header_row = 3
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=str(col_name))
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for r_idx, row in enumerate(rows, start=header_row + 1):
            for c_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=r_idx, column=c_idx, value=row.get(col_name))

        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        for col_idx, col_name in enumerate(columns, start=1):
            max_len = max([len(str(col_name))] + [len(str(r.get(col_name, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    spec = {
        "title": "Test Report",
        "sections": [{"heading": "Vendors", "data": [{"vendor": "Infosys", "spend": 100}, {"vendor": "Wipro", "spend": 80}]}],
    }
    data = build_excel_report(spec)
    assert data[:2] == b"PK", "not a valid xlsx (zip) file"
    print("excel OK —", len(data), "bytes")
