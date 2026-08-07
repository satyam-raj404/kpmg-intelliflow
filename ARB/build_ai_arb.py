"""Build the AI-Based ARB from the filled IntelliSource ARB.

Base = 'ARB Solution Review - Intellisource.docx.docx' (fully filled).
Adds the AI-specific ARB sections (AI NFRs, Data Governance, Operation &
Monitoring, AI System Impact Assessment), enriches AI-related tables, and
swaps in the 5 AI-updated architecture diagrams.

Output = 'ARB Solution Review - AI Based solution.docx.docx'
"""
import copy
import os
import shutil
import zipfile

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "ARB Solution Review - Intellisource.docx.docx")
OUT = os.path.join(HERE, "ARB Solution Review - AI Based solution.docx.docx")
DIAG = os.path.join(HERE, "Diagrams")

NAVY = RGBColor(0x00, 0x33, 0x8D)

# Media parts (in the base docx) → new AI-updated diagram PNGs to swap in.
IMAGE_SWAPS = {
    "word/media/image3.png":  os.path.join(DIAG, "D01_BusinessProcess_ToBe.png"),
    "word/media/image7.png":  os.path.join(DIAG, "D05_AppArch_ConceptualLogical.png"),
    "word/media/image11.png": os.path.join(DIAG, "D06_AppArch_TechnologyStack.png"),
    "word/media/image12.png": os.path.join(DIAG, "D07_IntegrationArch_SolutionContext.png"),
    "word/media/image13.png": os.path.join(DIAG, "D08_TechArch_AzureDeployment.png"),
}


# ── cell / table helpers ─────────────────────────────────────────────────────────

def set_cell(cell, text, bold=False, color=None):
    """Replace a cell's content with one or more paragraphs (split on \\n)."""
    cell.text = ""
    lines = text.split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        para = p if i == 0 else cell.add_paragraph()
        run = para.add_run(line)
        run.font.size = Pt(9)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color


def group_row(table, text):
    """Add a full-width-ish bold group header row (spanning first cell)."""
    row = table.add_row()
    set_cell(row.cells[0], text, bold=True, color=NAVY)
    # merge across the row for a clean band
    if len(row.cells) > 1:
        merged = row.cells[0]
        for c in row.cells[1:]:
            merged = merged.merge(c)
    return row


def kv_row(table, k, v):
    row = table.add_row()
    set_cell(row.cells[0], k, bold=True)
    set_cell(row.cells[1], v)
    return row


def make_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = NAVY
    return p


def apply_borders(new_tbl, template_tbl):
    """Clone the <w:tblBorders> from an existing bordered table."""
    src_pr = template_tbl._tbl.tblPr
    src_borders = src_pr.find(qn("w:tblBorders")) if src_pr is not None else None
    if src_borders is None:
        # fallback: build a plain single-line grid
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "999999")
            borders.append(e)
    else:
        borders = copy.deepcopy(src_borders)
    new_tbl._tbl.tblPr.append(borders)


def make_2col_table(doc, rows, template_tbl):
    """rows = list of ('group', title) | ('kv', k, v)."""
    t = doc.add_table(rows=0, cols=2)
    apply_borders(t, template_tbl)
    for r in rows:
        if r[0] == "group":
            group_row(t, r[1])
        else:
            kv_row(t, r[1], r[2])
    return t


def move_before(anchor_el, new_el):
    anchor_el.addprevious(new_el)


# ══════════════════════════════════════════════════════════════════════════════════
# 1. Copy base → output, open
# ══════════════════════════════════════════════════════════════════════════════════
shutil.copyfile(SRC, OUT)
doc = Document(OUT)
T = doc.tables  # 0-based; Table N == T[N-1]
BORDER_TMPL = T[9]  # NFR table — source of border formatting for new tables


# ══════════════════════════════════════════════════════════════════════════════════
# 2. Version table (Table 1) — add AI revision row
# ══════════════════════════════════════════════════════════════════════════════════
vt = T[0]
# row index 3 exists and is empty
vrow = vt.rows[3].cells
set_cell(vrow[0], "Draft 1.1")
set_cell(vrow[1], "04 Aug 2026")
set_cell(vrow[2], "AI solution inclusion — Ask IntelliSource (Azure OpenAI GPT-4o) governance, integration, and impact assessment added")
set_cell(vrow[3], "Satyam Barnwal")


# ══════════════════════════════════════════════════════════════════════════════════
# 3. Background — append AI note to Solution Name & Brief (Table 5, row 0, col 1)
# ══════════════════════════════════════════════════════════════════════════════════
bt = T[4]
brief_cell = bt.rows[0].cells[1]
brief_cell.add_paragraph(
    "The platform also includes Ask IntelliSource — an AI assistant that answers "
    "natural-language procurement questions grounded on the live database via read-only "
    "tool-calling (Azure OpenAI GPT-4o, KPMG tenant; pilot on OpenRouter GPT-4o)."
).runs[0].font.size = Pt(9)


# ══════════════════════════════════════════════════════════════════════════════════
# 4. Functional Requirements (Table 9) — enrich FR09
# ══════════════════════════════════════════════════════════════════════════════════
fr = T[8]
for row in fr.rows:
    if row.cells[0].text.strip() == "FR09":
        set_cell(row.cells[2],
            "Natural-language query interface (Ask IntelliSource) powered by GPT-4o. "
            "Target hosting: Azure OpenAI in KPMG tenant over Private Link; pilot uses "
            "OpenRouter GPT-4o (external — to be retired before production). Agentic "
            "tool-calling loop with 6 READ-ONLY tools: query_database, get_kpis, "
            "find_document, get_anomalies, get_vendor_info, get_p2p_stage_summary. "
            "Always calls a tool before answering data questions; never fabricates "
            "numbers; no write access to SAP or PostgreSQL.")
        break


# ══════════════════════════════════════════════════════════════════════════════════
# 5. Insert AI testing-scenario paragraphs after the FR table
#    (anchor = the "*Non-Functional Requirements" heading paragraph)
# ══════════════════════════════════════════════════════════════════════════════════
nfr_heading = None
compliance_heading = None
for p in doc.paragraphs:
    tx = p.text.strip()
    if tx == "*Non-Functional Requirements" and nfr_heading is None:
        nfr_heading = p
    if tx.startswith("*Compliance and Standard") and compliance_heading is None:
        compliance_heading = p

testing_paras = [
    ("** In case of an AI-Based solution — Testing scenarios for Ask IntelliSource:", True),
    ("Worst-case scenarios tested: adversarial / prompt-injection strings embedded in "
     "uploaded procurement data; malformed or ambiguous natural-language questions; "
     "requests that attempt to modify data (must be refused — the assistant is read-only); "
     "over-large result sets (mitigated by LIMIT injection and result compression); "
     "LLM timeout and rate-limit (429) conditions; and concurrent chat sessions.", False),
    ("Simulation / stress-testing: Yes (planned before production go-live). Methodology — "
     "(a) curated question bank with known ground-truth answers for accuracy regression; "
     "(b) red-team prompt set for injection / jailbreak attempts; (c) concurrent-session "
     "load test against the /api/chat endpoint. Tools — pytest regression harness, scripted "
     "load generator, and manual red-team review.", False),
]
if nfr_heading is not None:
    for text, bold in testing_paras:
        np = doc.add_paragraph()
        run = np.add_run(text)
        run.font.size = Pt(9.5)
        run.font.bold = bold
        move_before(nfr_heading._p, np._p)


# ══════════════════════════════════════════════════════════════════════════════════
# 6. Insert the four AI-specific sections BEFORE the Compliance heading
# ══════════════════════════════════════════════════════════════════════════════════
# Build them, then move each element before the compliance anchor, in order.

anchor = compliance_heading._p

def emit_heading(text):
    p = make_heading(doc, text)
    move_before(anchor, p._p)

def emit_table(rows):
    t = make_2col_table(doc, rows, BORDER_TMPL)
    move_before(anchor, t._tbl)
    # a spacer paragraph after each table
    sp = doc.add_paragraph()
    move_before(anchor, sp._p)

# --- 6a. AI Non-Functional Requirements ---
emit_heading("Non-Functional Requirements (only for AI-Based solutions)")
emit_table([
    ("group", "Human agency and oversight"),
    ("kv", "Level of autonomy",
     "No autonomous action. Ask IntelliSource is advisory and read-only — it answers "
     "questions and never writes to SAP or PostgreSQL, and never initiates a transaction. "
     "A human decision-maker is always in the loop; the assistant informs, humans act. "
     "Threshold for human intervention = all decisions."),
    ("kv", "Need for human reviewers",
     "Yes. All outputs are informational. Answers are grounded on the same live data the "
     "user can already see in dashboards and are traceable to the tool/query used "
     "(tools_used is returned), so users can validate. No AI output auto-executes a business action."),
    ("kv", "Possible dangers of removing human oversight",
     "Fully automated decision-making is NOT enabled and is out of scope. If removed, "
     "incorrect figures (mis-generated SQL or model hallucination) could mislead procurement "
     "decisions. Mitigations: mandatory tool-grounding (must call a tool before answering data "
     "questions), read-only SQL guard, numeric answers traceable to source, temperature 0.1 for determinism."),
    ("group", "Reliability / Technical Robustness"),
    ("kv", "Predictable failures in the model and mitigations",
     "Robustness testing planned before production (Transition-1). Known failure modes and "
     "mitigations: (a) invalid SQL → read-only guard + LIMIT injection + error surfaced, not "
     "fabricated; (b) hallucinated numbers → 'never invent numbers, always call a tool' system "
     "rule; (c) LLM timeout / rate-limit → retry with back-off + graceful fallback message; "
     "(d) prompt injection via data values → parameterised structured tools + read-only DB role. "
     "Adversarial and stress testing via a curated question bank, red-team prompts, and concurrent-session load."),
    ("group", "Transparency and Explainability"),
    ("kv", "Comprehensibility of AI decisions",
     "Not a black box in operation. The assistant exposes which tools / queries it used and its "
     "figures are cross-checkable against dashboards. The underlying GPT-4o model is itself "
     "opaque; this is justified because the model is used only as a language / reasoning layer over "
     "deterministic, auditable tool outputs — the authoritative data comes from governed read-only SQL, not the model."),
    ("kv", "Complexity of AI systems",
     "Moderate. Agentic tool-calling loop (max 8 iterations, guidance of ≤3 tool calls per answer). "
     "Complexity is justified: multi-table P2P questions require reasoning across "
     "PR → PO → GRN → Invoice → Payment, which simple keyword search cannot answer."),
    ("group", "Diversity, non-discrimination and fairness"),
    ("kv", "Demographics",
     "No impact on demographic groups. Data is internal KPMG procurement transactions (vendors, "
     "POs, spend); no personal or protected attributes of individuals are analysed and no profiling "
     "of persons occurs. Residual bias risk is limited to vendor treatment and is mitigated because "
     "answers reflect only factual recorded data, not model opinion."),
    ("kv", "Societal well-being",
     "Indirect positive impact — improves procurement compliance, MSME payment monitoring, and spend "
     "transparency. No adverse societal impact identified."),
    ("group", "Involvement of Individuals"),
    ("kv", "Involvement of affected individuals",
     "Affected individuals = KPMG internal procurement, finance, and leadership users, identified in "
     "the stakeholder analysis. Feedback via in-app feedback and the service desk. Users retain full "
     "control — the AI only informs; humans decide and act."),
    ("group", "Environmental Wellbeing"),
    ("kv", "Energy efficiency in model training",
     "Yes — a pre-trained hosted foundation model (Azure OpenAI GPT-4o) is used. No training or "
     "fine-tuning from scratch; zero training compute. Inference only, on demand."),
    ("kv", "Optimizing resource utilization",
     "Managed / serverless inference (Azure OpenAI) — no idle GPU owned. Azure Container Apps scale "
     "to zero; LLM calls are made only on a user query."),
    ("kv", "Sustainable data management",
     "No vector database or embeddings index is maintained (tool-calling on live relational data, not "
     "a RAG store) → minimal additional storage. No redundant AI data retained beyond capped "
     "in-session history (≤20 messages)."),
    ("kv", "Eco-friendly infrastructure",
     "Hosted on Microsoft Azure (serverless Container Apps + managed Azure OpenAI). Azure is committed "
     "to 100% renewable energy and carbon-negative operations."),
])

# --- 6b. Data Governance ---
emit_heading("Data Governance (AI)")
emit_table([
    ("kv", "Primary data sources",
     "Yes — proprietary internal data only. Sources: the live PostgreSQL procurement database "
     "(po_dump, pr_dump, grn_dump, po_invoice_dump, invoice_dump, payment_dump, vendor_master, "
     "kpi_results, pr_po_grn_invoice, process_mining_events). No public data and no third-party "
     "knowledge base. The model receives only the specific rows returned by governed read-only tools "
     "plus the schema description — never the whole database."),
    ("kv", "Data versioning and update frequency",
     "The underlying data is refreshed on each CSV upload and KPIs are recomputed. The assistant "
     "always queries live data at question time, so answers reflect the current database state — there "
     "is no stale cached AI knowledge and no external knowledge source to version."),
    ("kv", "Self-correction / discarding of incorrect data",
     "The system prompt forces tool-grounding and forbids inventing numbers; if a tool errors, the "
     "error is returned rather than masked. A regex write-guard blocks INSERT/UPDATE/DELETE/DDL, and "
     "LIMIT + result compression prevent oversized or inconsistent context. Harmful / biased output "
     "filtering via Azure OpenAI content filters (target); answers are constrained to factual "
     "procurement data."),
    ("group", "Model Robustness"),
    ("kv", "Metrics for go-live",
     "Answer accuracy against a ground-truth question bank ≥ agreed threshold; tool-grounding rate = "
     "100% for data questions; zero write operations; P95 response-latency target; endpoint uptime target."),
    ("kv", "Success metrics",
     "Latency, accuracy, throughput, and system uptime, measured against the question bank and "
     "monitored in production."),
    ("kv", "Responsibilities for checking model performance",
     "Development team plus a KPMG Responsible-AI reviewer monitor accuracy and regression; IT "
     "Operations monitors availability and latency."),
    ("kv", "Model confidence",
     "GPT-4o does not emit calibrated probabilities; confidence is handled by design — the deterministic "
     "tool outputs are authoritative. Uncertain or unanswerable questions cause the assistant to ask "
     "for clarification or state it cannot answer rather than guessing; temperature 0.1 reduces variance."),
])

# --- 6c. Operation & Monitoring ---
emit_heading("Operation & Monitoring (AI)")
emit_table([
    ("kv", "Components to be monitored",
     "AI request / response logs, tool-call logs (tools_used), LLM latency and error rate, token "
     "usage / cost, database query performance, chat-endpoint uptime, and Azure OpenAI service health."),
    ("kv", "Training programs for AI users",
     "In-app guidance and tooltips explain that the assistant is advisory and that figures should be "
     "verified against dashboards; onboarding notes cover its limitations (read-only, data-grounded)."),
    ("kv", "Decision-making power of AI",
     "None autonomous — advisory only. The assistant never executes a business action."),
    ("kv", "Process to overrule AI decisions",
     "Inherent: because the AI never executes anything, the user simply disregards or verifies an "
     "answer. No action-override mechanism is required since no actions are taken."),
    ("kv", "Monitoring input and output quality",
     "Yes — prompts and answers are logged; periodic sampling reviews answers against source data; "
     "error responses are captured."),
    ("kv", "Detecting drift in input data",
     "The model is externally hosted and managed (Azure OpenAI); model drift is managed by the provider "
     "plus version pinning of the deployment. Data drift is not applicable (no fine-tuning). A "
     "question-bank regression is run on any model-version change to trigger re-validation."),
    ("kv", "Versioned code repository",
     "Yes — all AI code (chat_engine, chat_tools, prompts, tool schemas) is version-controlled in "
     "GitHub, and the model deployment version is pinned in configuration."),
    ("kv", "Maintenance and update schedule",
     "Model deployment reviewed on Azure OpenAI version updates; code follows the standard release "
     "cycle; prompt changes are change-controlled."),
    ("kv", "Responsibilities for updates",
     "Development team (code / prompts); IT Operations (Azure OpenAI deployment); Responsible-AI "
     "reviewer (sign-off on prompt or model changes)."),
    ("kv", "Software tests",
     "Yes — question-bank regression, read-only guard tests, and integration tests run before deploy."),
    ("kv", "Stability period before deployment",
     "A defined soak / monitoring period in QA precedes promotion to Production."),
    ("kv", "Reporting AI output concerns",
     "In-app feedback and the service desk route concerns to the development team and the Responsible-AI "
     "reviewer; an incident process covers material errors."),
])

# --- 6d. AI System Impact Assessment ---
emit_heading("AI — System Impact Assessment (AISIA)")
aisia = doc.add_table(rows=0, cols=2)
apply_borders(aisia, BORDER_TMPL)
aisia.add_row()
# header row proper (2 cols)
hdr = aisia.rows[0]
set_cell(hdr.cells[0], "Requirements ID", bold=True)
set_cell(hdr.cells[1], "Assessment", bold=True)
r = aisia.add_row()
set_cell(r.cells[0], "AISIA 01")
set_cell(r.cells[1],
    "System: Ask IntelliSource — natural-language P2P analytics assistant.\n"
    "Purpose: answer procurement questions grounded on live internal data.\n"
    "Autonomy: advisory and read-only; human-in-the-loop for every decision.\n"
    "Data: internal KPMG procurement transactions only; no personal or special-category "
    "data; no data is used to train the model.\n"
    "Model & hosting: Azure OpenAI GPT-4o in the KPMG tenant over Private Link (target); "
    "the pilot uses OpenRouter GPT-4o (external, US-hosted) which must be retired before "
    "production due to data-egress risk.\n"
    "Key risks: hallucination, incorrect SQL generation, prompt injection, and external data "
    "egress (pilot).\n"
    "Mitigations: mandatory tool-grounding, read-only database role, LIMIT and result caps, "
    "low temperature (0.1), Azure OpenAI content filters, and in-tenant hosting for production.\n"
    "Impact classification: Low-to-Medium — an advisory internal tool with no automated decisions "
    "and no PII.\n"
    "Residual risk: acceptable for the pilot; migration to Azure OpenAI plus robustness testing "
    "is required before production.")
move_before(anchor, aisia._tbl)
sp = doc.add_paragraph(); move_before(anchor, sp._p)


# ══════════════════════════════════════════════════════════════════════════════════
# 7. Integration table (Table 32) — replace NA rows with the Azure OpenAI integration
# ══════════════════════════════════════════════════════════════════════════════════
it = T[31]
# rows: 0 = merged header, 1 = subheaders, 2/3 = NA data rows
dr = it.rows[2].cells
set_cell(dr[0], "IntelliSource Platform (FastAPI backend)")
set_cell(dr[1], "Azure OpenAI — GPT-4o (KPMG tenant)")
set_cell(dr[2],
    "NL question + tool schemas sent over HTTPS/JSON via Private Link; model returns tool "
    "calls / final answer. Only query-scoped result rows and the schema are shared — no bulk "
    "data. Pilot: OpenRouter GPT-4o (external) — to be replaced before production.")
set_cell(dr[3], "New")
dr2 = it.rows[3].cells
set_cell(dr2[0], "Azure Active Directory")
set_cell(dr2[1], "IntelliSource Platform")
set_cell(dr2[2], "SSO / OAuth 2.0 identity and role tokens (Target Architecture).")
set_cell(dr2[3], "New")


# ══════════════════════════════════════════════════════════════════════════════════
# 8. Bill of Materials (Table 39) + Licensing (Table 40) — add Azure OpenAI
# ══════════════════════════════════════════════════════════════════════════════════
bom = T[38]
brow = bom.add_row().cells
set_cell(brow[0], "Azure OpenAI (GPT-4o)")
set_cell(brow[1], "Microsoft Azure")
set_cell(brow[2], "GPT-4o (deployment)")
set_cell(brow[3], "GPT-4o (deployment)")
set_cell(brow[4], "Managed cloud service")
set_cell(brow[5], "Azure OpenAI")

lic = T[39]
lrow = lic.add_row().cells
set_cell(lrow[0],
    "Azure OpenAI Service (GPT-4o) — hosted LLM inference for the Ask IntelliSource assistant. "
    "Consumption-based (per 1K tokens) under Azure Enterprise Agreement. Pilot used OpenRouter "
    "(pay-per-use, external) — to be retired before production.")
set_cell(lrow[1], "Consumption-based")
set_cell(lrow[2], "Per 1K tokens (Azure EA)")
set_cell(lrow[3], "KPMG")
set_cell(lrow[4], "N/A")
set_cell(lrow[5], "N/A")


# ══════════════════════════════════════════════════════════════════════════════════
# 9. Risks (Table 41, single column) — append AI risks
# ══════════════════════════════════════════════════════════════════════════════════
risks = T[40]
ai_risks = [
    "AI — Hallucination / incorrect SQL: the LLM may generate a wrong figure or query. "
    "Mitigated by mandatory tool-grounding, the read-only guard, the 'never invent numbers' "
    "rule, and temperature 0.1; robustness testing required before production.",
    "AI — Data egress: the pilot sends query results to the external OpenRouter / OpenAI API "
    "(US-hosted). Procurement data is KPMG Confidential — must migrate to Azure OpenAI (KPMG "
    "tenant, Private Link) before production.",
    "AI — Prompt injection: malicious strings embedded in uploaded data could attempt to "
    "manipulate the model. Mitigated by parameterised structured tools and a read-only database "
    "role — no tool can perform a write.",
]
for r in ai_risks:
    set_cell(risks.add_row().cells[0], r)


# ══════════════════════════════════════════════════════════════════════════════════
# 10. Team Dependencies (Table 42) — add Responsible AI office
# ══════════════════════════════════════════════════════════════════════════════════
team = T[41]
trow = team.add_row().cells
set_cell(trow[0], "KPMG Responsible AI / AI Governance")
set_cell(trow[1], "P")
set_cell(trow[2],
    "Review of Ask IntelliSource against KPMG AI policy; AI System Impact Assessment (AISIA) "
    "sign-off and Azure OpenAI data-handling review before production go-live.")


# ══════════════════════════════════════════════════════════════════════════════════
# 11. Save (python-docx), then swap diagram media blobs in the saved zip
# ══════════════════════════════════════════════════════════════════════════════════
doc.save(OUT)

# swap media parts
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in IMAGE_SWAPS:
            with open(IMAGE_SWAPS[item.filename], "rb") as f:
                data = f.read()
        zout.writestr(item, data)
os.replace(tmp, OUT)

print("Saved:", OUT)
print("Tables:", len(Document(OUT).tables))
