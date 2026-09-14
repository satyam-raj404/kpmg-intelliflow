"""Case study slide — ribbon-banner layout (Client Name & Problem Statement /
Engagement Details / Technology Stack Used / Expected Outcome & Deliverables),
matching the second reference template. Same grounded content as
build_case_study_pptx.py, different visual format.

Output: presentations/IntelliSource_Case_Study_Ribbon.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY   = RGBColor(0x0B, 0x1F, 0x45)
MED    = RGBColor(0x00, 0x5E, 0xB8)
LIGHT  = RGBColor(0x00, 0x9F, 0xDA)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLACK  = RGBColor(0x1A, 0x1A, 0x1A)
GRAY   = RGBColor(0x55, 0x55, 0x55)
ICON_BG = [RGBColor(0xE8, 0x8A, 0x2E), RGBColor(0xBF, 0xE9, 0xF7),
           RGBColor(0xD1, 0x4A, 0x2E), RGBColor(0x2E, 0x8B, 0x57)]

W, H = Inches(13.33), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
slide = prs.slides.add_slide(prs.slide_layouts[6])


def rect(x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def rrect(x, y, w, h, color, radius=0.5):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def circle(x, y, d, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.color.rgb = WHITE; s.line.width = Pt(2)
    s.shadow.inherit = False
    return s


def txt(x, y, w, h, text, size=12, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, italic=False, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return box


def para_block(x, y, w, h, lines, size=10.5, color=BLACK, bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold_first and i == 0
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def bullet_block(x, y, w, h, items, size=10, color=BLACK, sub_items=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(7)
        r = p.add_run()
        r.text = f"•  {item}"
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = "Calibri"
    return box


# ── Title ──────────────────────────────────────────────────────────────────
txt(Inches(0.45), Inches(0.3), Inches(12.4), Inches(0.7),
    "IntelliSource — Procurement Process Intelligence", size=28, bold=True, color=NAVY)

# ── Ribbon banner: 4 segments ─────────────────────────────────────────────
ribbon_y = Inches(1.25)
ribbon_h = Inches(1.0)
seg_w = (W - Inches(0.9)) / 4
seg_colors = [NAVY, MED, LIGHT, NAVY]
headers = ["Client Name &\nProblem Statement", "Engagement Details",
           "Technology Stack Used", "Expected Outcome &\nDeliverables"]

x = Inches(0.45)
for i in range(4):
    rrect(x, ribbon_y, seg_w, ribbon_h, seg_colors[i], radius=0.5)
    x += seg_w - Inches(0.35)  # slight overlap for a connected ribbon look

# headers + icon circles drawn after, on top
x = Inches(0.45)
for i, hdr in enumerate(headers):
    icon_d = Inches(0.62)
    icon_x = x + seg_w - Inches(0.75)
    icon_y = ribbon_y + (ribbon_h - icon_d) / 2
    circle(icon_x, icon_y, icon_d, ICON_BG[i])
    box = slide.shapes.add_textbox(x + Inches(0.2), ribbon_y, seg_w - Inches(1.1), ribbon_h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(0)
    for j, line in enumerate(hdr.split("\n")):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE
    x += seg_w - Inches(0.35)

# ── Column content beneath ribbon ────────────────────────────────────────
content_y = ribbon_y + ribbon_h + Inches(0.35)
content_h = Inches(4.6)
x = Inches(0.45)

# Col 1 — Client & problem
para_block(x, content_y, seg_w - Inches(0.3), Inches(0.35),
           ["The Financial Services Team"], size=13, bold_first=True, color=BLACK)
para_block(x, content_y + Inches(0.42), seg_w - Inches(0.3), content_h - Inches(0.42),
           ["An internal KPMG India team needed a unified way to track and analyse "
            "procurement KPIs across the full Procure-to-Pay lifecycle — with no "
            "systematic way to flag process anomalies or answer ad-hoc questions "
            "without manual SQL/Excel work."], size=10.5, color=GRAY)
x += seg_w - Inches(0.35)

# Col 2 — Engagement details
para_block(x, content_y, seg_w - Inches(0.3), Inches(0.55),
           ["IntelliSource — Procurement Process Intelligence Platform"], size=12,
           bold_first=True, color=BLACK)
para_block(x, content_y + Inches(0.62), seg_w - Inches(0.3), content_h - Inches(0.62),
           ["Built an end-to-end platform automating 40+ KPIs across PR, PO, GRN, "
            "invoice, payment, vendor, and budget data, with real-time rule-based "
            "anomaly detection and a natural-language AI analyst (“Ask IntelliSource”) "
            "grounded in live data."], size=10.5, color=GRAY)
x += seg_w - Inches(0.35)

# Col 3 — Tech stack (bullets)
bullet_block(x, content_y, seg_w - Inches(0.3), content_h, [
    "Backend — Python, FastAPI, PostgreSQL (psycopg3)",
    "Frontend — React 19, Vite, TanStack Router, Tailwind CSS, Recharts",
    "AI — agentic LLM cascade via OpenRouter, local Ollama fallback",
    "Reporting — python-pptx, ReportLab, XlsxWriter, Matplotlib",
], size=10, color=BLACK)
x += seg_w - Inches(0.35)

# Col 4 — Outcomes & deliverables (bullets + value addition)
bullet_block(x, content_y, seg_w - Inches(0.3), Inches(2.6), [
    "Automated 40+ KPIs across procurement, delivery, finance, and compliance layers",
    "Real-time detection across 9 rule-based anomaly categories (maverick buys, duplicate invoices, 3-way mismatches, and more)",
    "Natural-language analyst answering ad-hoc procurement questions, grounded in live data with citations",
    "Role-based dashboards for Procurement Manager, Finance, Compliance, CXO, and more",
], size=9.5, color=BLACK)
txt(x, content_y + Inches(2.75), seg_w - Inches(0.3), Inches(0.3),
    "Value Addition:", size=10.5, bold=True, color=BLACK)
bullet_block(x, content_y + Inches(3.08), seg_w - Inches(0.3), Inches(1.3), [
    "Single source of truth, replacing manual Excel-based reporting",
    "Significant reduction in manual KPI-computation effort",
], size=9.5, color=BLACK)

# ── Bottom "Illustrative Deliverables" bar ─────────────────────────────────
bar_y = H - Inches(0.55)
bar_h = Inches(0.32)
rrect(Inches(0.45), bar_y, W - Inches(0.9), bar_h, NAVY, radius=0.5)
txt(Inches(0.45), bar_y, W - Inches(0.9), bar_h, "Illustrative Deliverables",
    size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for cx in (Inches(0.45), W - Inches(0.45)):
    circle(cx - Inches(0.08), bar_y - Inches(0.08), Inches(0.16), NAVY)
    circle(cx - Inches(0.08), bar_y + bar_h - Inches(0.08), Inches(0.16), NAVY)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "presentations",
                    "IntelliSource_Case_Study_Ribbon.pptx")
prs.save(out)
print(f"Saved: {os.path.abspath(out)}")
