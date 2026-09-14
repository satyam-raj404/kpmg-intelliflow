"""Case study slide — IntelliSource, built for The Financial Services Team
(internal), matching the reference KPMG case-study template layout: title
bar, light-blue intro banner, then four gradient columns (Key Challenges /
Tools Utilized / Our Approach / Our Value Add).

All content is grounded in the actual codebase (backend/services/*.py,
requirements.txt, package.json) — not fabricated. Measurable outcomes are
scope facts (KPI count, anomaly categories) rather than invented metrics,
since this hasn't been deployed to production yet and has no usage telemetry.

Output: presentations/IntelliSource_Case_Study.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── KPMG-style palette (matches scripts/build_presentation_pptx.py) ──────────
NAVY    = RGBColor(0x00, 0x33, 0x8D)
MED     = RGBColor(0x00, 0x5E, 0xB8)
LIGHT   = RGBColor(0x00, 0x91, 0xDA)
GOLD    = RGBColor(0x8F, 0x73, 0x26)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x05, 0x18, 0x35)
INTRO_BG = RGBColor(0xBF, 0xE9, 0xF7)  # light cyan intro banner
PURPLE  = RGBColor(0x47, 0x0A, 0x68)   # 4th column
TEXT_MUTED = RGBColor(0x33, 0x33, 0x33)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def rect(x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = color
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rrect(x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(x, y, w, h, text, size=12, bold=False, color=DARK, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, italic=False, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return box


def bullets(x, y, w, h, items, size=10.5, color=WHITE):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = f"•  {item}"
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


# ── Title bar ──────────────────────────────────────────────────────────────
rect(0, 0, W, Inches(1.15), WHITE)
txt(Inches(0.4), Inches(0.28), Inches(6), Inches(0.6),
    "Case Study — IntelliSource", size=30, bold=True, color=NAVY)
txt(Inches(0.4), Inches(0.82), Inches(8), Inches(0.3),
    "AI-Powered Procurement Process Intelligence", size=13, color=MED, italic=True)

# tech "logo row" (text chips — no external logo assets available)
chips = ["FastAPI", "PostgreSQL", "React", "LLM Agent"]
cx = Inches(9.3)
for c in chips:
    w = Inches(0.95)
    rrect(cx, Inches(0.4), w, Inches(0.4), NAVY)
    txt(cx, Inches(0.4), w, Inches(0.4), c, size=9, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += w + Inches(0.1)

# ── Intro banner ───────────────────────────────────────────────────────────
banner_y = Inches(1.15)
banner_h = Inches(1.45)
rect(0, banner_y, W, banner_h, INTRO_BG)
txt(Inches(0.4), banner_y + Inches(0.1), Inches(12.5), Inches(0.65),
    "The Financial Services Team (India) needed a unified way to track, automate, and "
    "analyse procurement KPIs across the full Procure-to-Pay lifecycle — with anomaly "
    "detection and natural-language analytics on top of live data.",
    size=13, bold=True, color=DARK)
txt(Inches(0.4), banner_y + Inches(0.72), Inches(3), Inches(0.3), "Our scope included:",
    size=11, bold=True, color=DARK)
bullets(Inches(0.4), banner_y + Inches(1.0), Inches(12.5), Inches(0.4), [
    "40+ KPI automation across PR/PO/GRN/Invoice/Payment, vendor, budget, and compliance data",
    "Real-time rule-based anomaly detection and a grounded AI analyst for ad-hoc procurement questions",
], size=10.5, color=DARK)

# ── Four columns ───────────────────────────────────────────────────────────
col_y = banner_y + banner_h + Inches(0.75)
col_h = H - col_y - Inches(0.55)
gap = Inches(0.15)
col_w = (W - Inches(0.8) - gap * 3) / 4
col_x0 = Inches(0.4)

columns = [
    ("Key Challenges", LIGHT, [
        "Procurement data fragmented across PR, PO, GRN, invoice, and payment stages — no unified view",
        "40+ KPIs computed manually across procurement, delivery, finance, and compliance — slow and error-prone",
        "No systematic way to flag process anomalies (maverick buys, duplicate invoices, mismatches) in real time",
        "Ad-hoc procurement questions required manual SQL/Excel work with slow turnaround",
    ]),
    ("Tools Utilized", MED, [
        "Backend — Python, FastAPI, PostgreSQL (psycopg3)",
        "Frontend — React 19, Vite, TanStack Router, Tailwind CSS, Recharts",
        "AI — agentic LLM cascade via OpenRouter, local Ollama fallback",
        "Reporting — python-pptx, ReportLab, XlsxWriter, Matplotlib",
    ]),
    ("Our Approach", NAVY, [
        "Built a unified P2P data model spanning PR, PO, GRN, invoice, payment, vendor, and budget data",
        "Automated 40+ KPIs across procurement, delivery, finance, and compliance layers",
        "Implemented rule-based process-mining anomaly detection across 9 categories, flagged in real time",
        "Built “Ask IntelliSource” — a grounded, citation-backed natural-language analyst with live SQL tool access",
        "Delivered role-based dashboards (Procurement Manager, Finance, Compliance, CXO, and more)",
    ]),
    ("Our Value Add", PURPLE, [
        "Single source of truth for procurement KPIs, replacing manual Excel-based reporting",
        "Real-time visibility into 9 categories of process anomalies for faster intervention",
        "Natural-language access to procurement data — no SQL or BI tooling expertise required",
        "Extensible, provider-agnostic AI architecture — swap LLM models/providers via config, no code changes",
    ]),
]

x = col_x0
for title, color, items in columns:
    card = rrect(x, col_y, col_w, col_h, color)
    txt(x + Inches(0.18), col_y + Inches(0.15), col_w - Inches(0.36), Inches(0.4),
        title, size=13, bold=True, color=WHITE)
    rect(x + Inches(0.18), col_y + Inches(0.58), col_w - Inches(0.36), Pt(1.5), WHITE)
    bullets(x + Inches(0.18), col_y + Inches(0.72), col_w - Inches(0.36),
            col_h - Inches(0.9), items, size=9.5, color=WHITE)
    x += col_w + gap

# ── Footer ─────────────────────────────────────────────────────────────────
txt(Inches(0.4), H - Inches(0.42), Inches(9), Inches(0.3),
    "Internal — The Financial Services Team, KPMG India  |  Not yet deployed to production  |  "
    "Engagement owner: [TBD]",
    size=8, italic=True, color=TEXT_MUTED)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "presentations", "IntelliSource_Case_Study.pptx")
prs.save(out)
print(f"Saved: {os.path.abspath(out)}")
